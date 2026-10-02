from __future__ import annotations

import json
import re
from uuid import uuid4
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import fitz
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.constants import UserAccessPermissions as UAP

from .models import ConversionResult, PdfInfo
from .output_safety import safe_outputs
from .page_ranges import parse_page_selection, parse_split_ranges


class PdfToolkitError(Exception):
    """Raised when an operation on a PDF cannot be completed."""


@dataclass(slots=True)
class ProtectOptions:
    allow_print: bool = True
    allow_copy: bool = True
    allow_modify: bool = False
    allow_annotate: bool = True


class PdfToolkit:
    def __init__(self, output_root: Path) -> None:
        self.output_root = output_root
        self.output_root.mkdir(parents=True, exist_ok=True)

    def inspect(self, pdf_path: Path, password: str | None = None) -> PdfInfo:
        pdf_path = pdf_path.resolve()
        reader = PdfReader(str(pdf_path))
        if reader.is_encrypted:
            if not password:
                raise PdfToolkitError("This PDF is password-protected.")
            if reader.decrypt(password) == 0:
                raise PdfToolkitError("Invalid password for encrypted PDF.")

        metadata = reader.metadata or {}
        permissions = {
            "print": True,
            "copy": True,
            "modify": True,
            "annotate": True,
        }

        return PdfInfo(
            path=pdf_path,
            page_count=len(reader.pages),
            encrypted=reader.is_encrypted,
            title=str(metadata.get("/Title", "")),
            author=str(metadata.get("/Author", "")),
            subject=str(metadata.get("/Subject", "")),
            creator=str(metadata.get("/Creator", "")),
            producer=str(metadata.get("/Producer", "")),
            permissions=permissions,
        )

    @safe_outputs
    def merge(self, pdf_paths: Iterable[Path], output_path: Path) -> Path:
        paths = [Path(p).resolve() for p in pdf_paths]
        if len(paths) < 2:
            raise PdfToolkitError("Choose at least two PDF files to merge.")

        writer = PdfWriter()
        for path in paths:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise PdfToolkitError(f"'{path.name}' is encrypted. Unlock it first.")
            for page in reader.pages:
                writer.add_page(page)

        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("wb") as handle:
            writer.write(handle)
        return output_path

    @safe_outputs
    def split_by_ranges(self, pdf_path: Path, range_text: str, output_dir: Path) -> list[Path]:
        source = Path(pdf_path).resolve()
        reader = PdfReader(str(source))
        if reader.is_encrypted:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before split.")
        ranges = parse_split_ranges(range_text, len(reader.pages))

        output_dir = output_dir.resolve()
        output_dir.mkdir(parents=True, exist_ok=True)

        outputs: list[Path] = []
        base = source.stem
        for start, end in ranges:
            writer = PdfWriter()
            for index in range(start, end + 1):
                writer.add_page(reader.pages[index])
            name = f"{base}_pages_{start + 1}-{end + 1}.pdf"
            out = output_dir / name
            with out.open("wb") as handle:
                writer.write(handle)
            outputs.append(out)
        return outputs

    @safe_outputs
    def split_every(self, pdf_path: Path, pages_per_file: int, output_dir: Path) -> list[Path]:
        if pages_per_file < 1:
            raise PdfToolkitError("Pages per split file must be at least 1.")

        source = Path(pdf_path).resolve()
        reader = PdfReader(str(source))
        if reader.is_encrypted:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before split.")

        output_dir = output_dir.resolve()
        output_dir.mkdir(parents=True, exist_ok=True)

        outputs: list[Path] = []
        base = source.stem
        total = len(reader.pages)
        part = 1
        for start in range(0, total, pages_per_file):
            end = min(start + pages_per_file, total)
            writer = PdfWriter()
            for index in range(start, end):
                writer.add_page(reader.pages[index])
            out = output_dir / f"{base}_part_{part:03d}.pdf"
            with out.open("wb") as handle:
                writer.write(handle)
            outputs.append(out)
            part += 1
        return outputs

    def _select_pages(self, source: Path, order: list[int], output_path: Path) -> Path:
        with fitz.open(str(source)) as doc:
            if doc.needs_pass:
                raise PdfToolkitError("Unlock this PDF before changing pages.")
            if not order:
                raise PdfToolkitError("PDF must retain at least one page.")
            toc = doc.get_toc()
            first_destination = {}
            seen_pages = set()
            independent_order = []
            for new_index, old_index in enumerate(order):
                first_destination.setdefault(old_index, new_index)
                if old_index in seen_pages:
                    # select() otherwise aliases duplicate page xrefs, making
                    # bookmark destinations resolve to the last duplicate.
                    doc.fullcopy_page(old_index)
                    independent_order.append(doc.page_count - 1)
                else:
                    independent_order.append(old_index)
                    seen_pages.add(old_index)
            doc.select(independent_order)
            mapped = []
            for level, title, page_number in toc:
                target = first_destination.get(page_number - 1)
                if target is not None:
                    # A removed parent cannot leave an invalid level jump.
                    level = min(level, (mapped[-1][0] + 1) if mapped else 1)
                    mapped.append([level, title, target + 1])
            doc.set_toc(mapped)
            doc.save(str(output_path), garbage=3, deflate=True)
        return output_path

    @safe_outputs
    def extract_pages(self, pdf_path: Path, selection: str, output_path: Path) -> Path:
        with fitz.open(str(pdf_path)) as doc:
            order = parse_page_selection(selection, doc.page_count)
        return self._select_pages(pdf_path, order, output_path)

    @safe_outputs
    def delete_pages(self, pdf_path: Path, selection: str, output_path: Path) -> Path:
        with fitz.open(str(pdf_path)) as doc:
            deleted = set(parse_page_selection(selection, doc.page_count))
            order = [i for i in range(doc.page_count) if i not in deleted]
        return self._select_pages(pdf_path, order, output_path)

    @safe_outputs
    def reorder_pages(self, pdf_path: Path, order_text: str, output_path: Path) -> Path:
        with fitz.open(str(pdf_path)) as doc:
            order = self._parse_order_with_duplicates(order_text, doc.page_count)
        return self._select_pages(pdf_path, order, output_path)

    @safe_outputs
    def rotate_pages(self, pdf_path: Path, selection: str, degrees: int, output_path: Path) -> Path:
        if degrees not in {90, 180, 270}:
            raise PdfToolkitError("Rotation must be 90, 180, or 270 degrees.")
        with fitz.open(str(pdf_path)) as doc:
            if doc.needs_pass:
                raise PdfToolkitError("Unlock this PDF before rotating pages.")
            for index in parse_page_selection(selection, doc.page_count):
                page = doc[index]
                page.set_rotation((page.rotation + degrees) % 360)
            doc.save(str(output_path), garbage=3, deflate=True)
        return output_path

    @safe_outputs
    def watermark_text(
        self,
        pdf_path: Path,
        text: str,
        output_path: Path,
        selection: str | None = None,
        opacity: float = 0.2,
    ) -> Path:
        if not text.strip():
            raise PdfToolkitError("Watermark text cannot be empty.")
        if opacity <= 0 or opacity > 1:
            raise PdfToolkitError("Opacity must be between 0 and 1.")

        source = Path(pdf_path).resolve()
        doc = fitz.open(str(source))
        if doc.needs_pass:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before watermarking.")

        page_indices = parse_page_selection(selection, doc.page_count)
        for page_index in page_indices:
            page = doc.load_page(page_index)
            rect = page.rect
            font_size = max(24, min(72, int(rect.width / 12)))
            # Repeat watermark text across the page so it remains visible after cropping.
            step = max(90, int(font_size * 1.8))
            x_start = int(rect.width * 0.05)
            for y in range(step, int(rect.height), step):
                page.insert_text(
                    fitz.Point(x_start, y),
                    text,
                    fontsize=font_size,
                    color=(0.14, 0.2, 0.35),
                    fill_opacity=opacity,
                    overlay=True,
                )

        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(
            str(output_path),
            garbage=4,
            clean=True,
            deflate=True,
        )
        doc.close()
        return output_path

    @safe_outputs
    def protect(self, pdf_path: Path, user_password: str, owner_password: str | None,
                options: ProtectOptions, output_path: Path) -> Path:
        if not user_password:
            raise PdfToolkitError("User password is required.")
        with fitz.open(str(pdf_path)) as doc:
            if doc.is_encrypted:
                raise PdfToolkitError("PDF is already encrypted. Unlock it first.")
            permissions = 0
            for allowed, flag in [(options.allow_print, fitz.PDF_PERM_PRINT),
                                  (options.allow_copy, fitz.PDF_PERM_COPY),
                                  (options.allow_modify, fitz.PDF_PERM_MODIFY),
                                  (options.allow_annotate, fitz.PDF_PERM_ANNOTATE)]:
                if allowed:
                    permissions |= flag
            doc.save(str(output_path), encryption=fitz.PDF_ENCRYPT_AES_256,
                     user_pw=user_password, owner_pw=owner_password or user_password,
                     permissions=permissions, deflate=True)
        return output_path

    @safe_outputs
    def unlock(self, pdf_path: Path, password: str, output_path: Path) -> Path:
        with fitz.open(str(pdf_path)) as doc:
            if not doc.is_encrypted:
                raise PdfToolkitError("PDF is not encrypted.")
            if not doc.authenticate(password):
                raise PdfToolkitError("Invalid password.")
            doc.save(str(output_path), encryption=fitz.PDF_ENCRYPT_NONE, deflate=True)
        return output_path

    @safe_outputs
    def compress(self, pdf_path: Path, output_path: Path) -> Path:
        source = Path(pdf_path).resolve()
        doc = fitz.open(str(source))
        if doc.needs_pass:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before optimization.")

        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(
            str(output_path),
            garbage=4,
            clean=True,
            deflate=True,
            use_objstms=1,
        )
        doc.close()
        return output_path

    @safe_outputs
    def annotate_text_matches(
        self,
        pdf_path: Path,
        query: str,
        selection: str,
        style: str,
        output_path: Path,
    ) -> Path:
        source = Path(pdf_path).resolve()
        if not query.strip():
            raise PdfToolkitError("Annotation query cannot be empty.")

        style_name = style.strip().lower()
        if style_name not in {"highlight", "underline", "strikeout"}:
            raise PdfToolkitError("Annotation style must be highlight, underline, or strikeout.")

        doc = fitz.open(str(source))
        if doc.needs_pass:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before annotation.")

        page_indices = parse_page_selection(selection, doc.page_count)
        hit_count = 0
        try:
            for page_index in page_indices:
                page = doc.load_page(page_index)
                for rect in page.search_for(query):
                    if style_name == "highlight":
                        annot = page.add_highlight_annot(rect)
                        annot.set_colors(stroke=(1.0, 0.91, 0.33))
                    elif style_name == "underline":
                        annot = page.add_underline_annot(rect)
                        annot.set_colors(stroke=(0.19, 0.42, 0.89))
                    else:
                        annot = page.add_strikeout_annot(rect)
                        annot.set_colors(stroke=(0.85, 0.27, 0.21))
                    annot.update()
                    hit_count += 1

            if hit_count == 0:
                raise PdfToolkitError("No text matches found for annotation.")

            output_path = output_path.resolve()
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(
                str(output_path),
                garbage=4,
                clean=True,
                deflate=True,
                use_objstms=1,
            )
            return output_path
        finally:
            doc.close()

    @safe_outputs
    def redact_text_matches(
        self,
        pdf_path: Path,
        query: str,
        selection: str,
        output_path: Path,
        fill: tuple[float, float, float] = (0.0, 0.0, 0.0),
    ) -> Path:
        source = Path(pdf_path).resolve()
        if not query.strip():
            raise PdfToolkitError("Redaction query cannot be empty.")

        doc = fitz.open(str(source))
        if doc.needs_pass:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before redaction.")

        page_indices = parse_page_selection(selection, doc.page_count)
        hit_count = 0
        try:
            for page_index in page_indices:
                page = doc.load_page(page_index)
                rects = page.search_for(query)
                if not rects:
                    continue
                for rect in rects:
                    page.add_redact_annot(rect, fill=fill)
                    hit_count += 1
                page.apply_redactions()

            if hit_count == 0:
                raise PdfToolkitError("No text matches found for redaction.")

            output_path = output_path.resolve()
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(
                str(output_path),
                garbage=4,
                clean=True,
                deflate=True,
                use_objstms=1,
            )
            return output_path
        finally:
            doc.close()

    @safe_outputs
    def stamp_image(
        self,
        pdf_path: Path,
        image_path: Path,
        selection: str,
        output_path: Path,
        *,
        anchor: str = "bottom-right",
        width_ratio: float = 0.22,
        margin: float = 24.0,
    ) -> Path:
        source = Path(pdf_path).resolve()
        stamp = Path(image_path).resolve()
        if not stamp.exists():
            raise PdfToolkitError(f"Stamp image not found: {stamp}")

        width_ratio = max(0.05, min(0.9, float(width_ratio)))
        margin = max(0.0, float(margin))
        anchor_name = anchor.strip().lower()
        if anchor_name not in {"top-left", "top-right", "bottom-left", "bottom-right", "center"}:
            raise PdfToolkitError("Stamp anchor must be top-left, top-right, bottom-left, bottom-right, or center.")

        doc = fitz.open(str(source))
        if doc.needs_pass:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before stamping.")

        try:
            try:
                stamp_pix = fitz.Pixmap(str(stamp))
            except Exception as exc:  # pragma: no cover - defensive for unsupported files.
                raise PdfToolkitError(f"Could not load stamp image: {stamp.name}") from exc

            if stamp_pix.width <= 0 or stamp_pix.height <= 0:
                raise PdfToolkitError("Stamp image has invalid dimensions.")
            aspect = stamp_pix.height / stamp_pix.width
            page_indices = parse_page_selection(selection, doc.page_count)

            for page_index in page_indices:
                page = doc.load_page(page_index)
                rect = page.rect
                target_width = min(rect.width * width_ratio, rect.width - margin * 2)
                target_width = max(36.0, target_width)
                target_height = target_width * aspect

                if target_height > rect.height - margin * 2:
                    target_height = max(28.0, rect.height - margin * 2)
                    target_width = target_height / max(0.001, aspect)

                x0 = rect.x0 + margin
                y0 = rect.y0 + margin
                if anchor_name in {"top-right", "bottom-right"}:
                    x0 = rect.x1 - margin - target_width
                if anchor_name in {"bottom-left", "bottom-right"}:
                    y0 = rect.y1 - margin - target_height
                if anchor_name == "center":
                    x0 = rect.x0 + (rect.width - target_width) / 2
                    y0 = rect.y0 + (rect.height - target_height) / 2

                target = fitz.Rect(x0, y0, x0 + target_width, y0 + target_height)
                page.insert_image(target, filename=str(stamp), keep_proportion=True, overlay=True)

            output_path = output_path.resolve()
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(
                str(output_path),
                garbage=4,
                clean=True,
                deflate=True,
                use_objstms=1,
            )
            return output_path
        finally:
            doc.close()

    @safe_outputs
    def convert(self, pdf_path: Path, target: str, output_dir: Path) -> ConversionResult:
        source = Path(pdf_path).resolve()
        doc = fitz.open(str(source))
        if doc.needs_pass:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before conversion.")

        target = target.lower().strip()
        output_dir = output_dir.resolve()
        output_dir.mkdir(parents=True, exist_ok=True)

        if target in {"txt", "md"}:
            text = self.extract_reflow_text(source)
            suffix = "txt" if target == "txt" else "md"
            out = output_dir / f"{source.stem}.{suffix}"
            if target == "txt":
                out.write_text(text, encoding="utf-8")
            else:
                markdown = self._as_markdown(text)
                out.write_text(markdown, encoding="utf-8")
            doc.close()
            return ConversionResult(target=target, outputs=[out])

        if target == "json":
            pages_payload: list[dict[str, object]] = []
            for idx, page in enumerate(doc, start=1):
                blocks = page.get_text("blocks")
                normalized_blocks: list[dict[str, object]] = []
                for block in blocks:
                    normalized_blocks.append(
                        {
                            "x0": float(block[0]),
                            "y0": float(block[1]),
                            "x1": float(block[2]),
                            "y1": float(block[3]),
                            "text": self._clean_page_text(str(block[4])),
                        }
                    )
                pages_payload.append(
                    {
                        "page": idx,
                        "text": self._clean_page_text(page.get_text("text")),
                        "blocks": normalized_blocks,
                    }
                )
            payload = {
                "source": str(source),
                "pages": pages_payload,
            }
            out = output_dir / f"{source.stem}.json"
            out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
            doc.close()
            return ConversionResult(target=target, outputs=[out])

        if target == "rtf":
            text = self.extract_reflow_text(source)
            out = output_dir / f"{source.stem}.rtf"
            out.write_text(self._as_rtf(text), encoding="utf-8")
            doc.close()
            return ConversionResult(target=target, outputs=[out])

        if target == "html":
            pages: list[str] = []
            for idx, page in enumerate(doc, start=1):
                body = page.get_text("html")
                pages.append(f"<section><h2>Page {idx}</h2>{body}</section>")
            html = (
                "<!doctype html><html><head><meta charset='utf-8'>"
                "<style>body{font-family:Segoe UI,Tahoma,sans-serif;padding:24px;}"
                "section{margin-bottom:40px;}h2{color:#1f3a5f;}</style>"
                "</head><body>"
                + "\n".join(pages)
                + "</body></html>"
            )
            out = output_dir / f"{source.stem}.html"
            out.write_text(html, encoding="utf-8")
            doc.close()
            return ConversionResult(target=target, outputs=[out])

        if target == "docx":
            document = Document()
            document.add_heading(source.stem, level=1)
            for idx, page in enumerate(doc, start=1):
                document.add_heading(f"Page {idx}", level=2)
                text = self._clean_page_text(page.get_text("text"))
                document.add_paragraph(text if text else "[No text detected on this page]")
            out = output_dir / f"{source.stem}.docx"
            document.save(str(out))
            doc.close()
            return ConversionResult(target=target, outputs=[out])

        if target in {"png", "jpg"}:
            outputs: list[Path] = []
            for idx, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                ext = "jpg" if target == "jpg" else "png"
                out = output_dir / f"{source.stem}_page_{idx:03d}.{ext}"
                pix.save(str(out))
                outputs.append(out)
            doc.close()
            return ConversionResult(target=target, outputs=outputs)

        doc.close()
        raise PdfToolkitError(
            "Unsupported conversion target. Supported: docx, txt, md, html, json, rtf, png, jpg."
        )

    @safe_outputs
    def convert_to_pdf(self, input_paths: Iterable[Path], output_path: Path) -> Path:
        sources = [Path(path).resolve() for path in input_paths]
        if not sources:
            raise PdfToolkitError("At least one source file is required for convert-to-PDF.")

        out_doc = fitz.open()
        try:
            for source in sources:
                if not source.exists():
                    raise PdfToolkitError(f"File not found: {source}")
                suffix = source.suffix.lower()

                if suffix == ".pdf":
                    src_doc = fitz.open(str(source))
                    if src_doc.needs_pass:
                        src_doc.close()
                        raise PdfToolkitError(f"Encrypted PDF not supported in convert-to-PDF: {source.name}")
                    out_doc.insert_pdf(src_doc)
                    src_doc.close()
                    continue

                if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
                    self._append_image_as_pdf_pages(out_doc, source)
                    continue

                if suffix in {".txt", ".md", ".rtf"}:
                    text = source.read_text(encoding="utf-8", errors="ignore")
                    self._append_text_as_pdf_pages(out_doc, text, source.name)
                    continue

                if suffix in {".html", ".htm"}:
                    html = source.read_text(encoding="utf-8", errors="ignore")
                    plain = self._strip_html(html)
                    self._append_text_as_pdf_pages(out_doc, plain, source.name)
                    continue

                if suffix == ".docx":
                    document = Document(str(source))
                    chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
                    text = "\n\n".join(chunks)
                    self._append_text_as_pdf_pages(out_doc, text, source.name)
                    continue

                raise PdfToolkitError(
                    f"Unsupported source format for convert-to-PDF: {source.suffix or source.name}"
                )

            output_path = output_path.resolve()
            output_path.parent.mkdir(parents=True, exist_ok=True)
            out_doc.save(
                str(output_path),
                garbage=4,
                clean=True,
                deflate=True,
                use_objstms=1,
            )
            return output_path
        finally:
            out_doc.close()

    def extract_reflow_text(self, pdf_path: Path) -> str:
        source = Path(pdf_path).resolve()
        doc = fitz.open(str(source))
        if doc.needs_pass:
            raise PdfToolkitError("Encrypted PDFs must be unlocked before text extraction.")

        pages: list[str] = []
        try:
            for page_number, page in enumerate(doc, start=1):
                blocks = page.get_text("blocks")
                block_texts: list[str] = []
                for block in sorted(blocks, key=lambda item: (item[1], item[0])):
                    text = self._clean_page_text(str(block[4]))
                    if text:
                        block_texts.append(text)
                joined = "\n\n".join(block_texts).strip()
                pages.append(f"Page {page_number}\n{joined}".strip())
            return "\n\n".join(pages).strip()
        finally:
            doc.close()

    def default_output_path(self, source_pdf: Path, operation: str, extension: str = "pdf") -> Path:
        now = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        name = f"{source_pdf.stem}_{operation}_{now}_{uuid4().hex[:8]}.{extension.lstrip('.')}"
        return self.output_root / name

    @staticmethod
    def _clean_page_text(raw: str) -> str:
        lines = [line.strip() for line in raw.replace("\r", "\n").split("\n")]
        compact = [line for line in lines if line]
        return "\n".join(compact)

    @staticmethod
    def _as_markdown(text: str) -> str:
        chunks = text.split("\n\n")
        rendered = []
        for chunk in chunks:
            line = chunk.strip()
            if not line:
                continue
            if line.startswith("Page "):
                rendered.append(f"## {line}")
                continue
            rendered.append(line)
        return "\n\n".join(rendered)

    @staticmethod
    def _as_rtf(text: str) -> str:
        escaped = text.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}")
        escaped = escaped.replace("\n", "\\par\n")
        return "{\\rtf1\\ansi\\deff0\n" + escaped + "\n}"

    @staticmethod
    def _strip_html(html: str) -> str:
        without_scripts = re.sub(r"(?is)<(script|style).*?>.*?</\\1>", " ", html)
        with_breaks = re.sub(r"(?i)<br\\s*/?>", "\n", without_scripts)
        with_breaks = re.sub(r"(?i)</p\\s*>", "\n\n", with_breaks)
        plain = re.sub(r"(?s)<[^>]+>", " ", with_breaks)
        plain = plain.replace("&nbsp;", " ").replace("&amp;", "&")
        return re.sub(r"[ \\t]+", " ", plain).replace("\r", "\n")

    def _append_image_as_pdf_pages(self, out_doc: fitz.Document, source: Path) -> None:
        img_doc = fitz.open(str(source))
        try:
            pdf_bytes = img_doc.convert_to_pdf()
            image_pdf = fitz.open("pdf", pdf_bytes)
            out_doc.insert_pdf(image_pdf)
            image_pdf.close()
        finally:
            img_doc.close()

    def _append_text_as_pdf_pages(self, out_doc: fitz.Document, text: str, title: str) -> None:
        page_rect = fitz.Rect(0, 0, 595, 842)  # A4 portrait in points.
        content_rect = fitz.Rect(50, 60, 545, 790)

        normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not normalized:
            normalized = "[No textual content]"

        wrapped_lines: list[str] = []
        for raw_line in normalized.split("\n"):
            line = raw_line.strip()
            if not line:
                wrapped_lines.append("")
                continue
            while len(line) > 96:
                split_at = line.rfind(" ", 0, 96)
                if split_at <= 0:
                    split_at = 96
                wrapped_lines.append(line[:split_at].rstrip())
                line = line[split_at:].lstrip()
            wrapped_lines.append(line)

        cursor = 0
        lines_per_page = 46
        first_page = True
        while cursor < len(wrapped_lines):
            page = out_doc.new_page(width=page_rect.width, height=page_rect.height)
            if first_page:
                page.insert_text(
                    fitz.Point(content_rect.x0, 38),
                    f"Source: {title}",
                    fontsize=10,
                    color=(0.25, 0.33, 0.46),
                )
                first_page = False

            chunk = "\n".join(wrapped_lines[cursor : cursor + lines_per_page]).strip() or "[Blank line]"
            page.insert_textbox(
                content_rect,
                chunk,
                fontsize=11,
                fontname="helv",
                color=(0.12, 0.16, 0.23),
                align=fitz.TEXT_ALIGN_LEFT,
                render_mode=0,
            )
            cursor += lines_per_page

    @staticmethod
    def _parse_order_with_duplicates(order_text: str, page_count: int) -> list[int]:
        if not order_text.strip():
            raise PdfToolkitError("Page order is required.")

        result: list[int] = []
        for token in order_text.split(","):
            item = token.strip()
            if not item:
                continue

            if "-" in item:
                start_text, end_text = item.split("-", 1)
                start = int(start_text)
                end = int(end_text)
                if start < 1 or end < 1 or start > page_count or end > page_count:
                    raise PdfToolkitError(f"Order range '{item}' is out of bounds.")
                step = 1 if end >= start else -1
                for page in range(start, end + step, step):
                    result.append(page - 1)
                continue

            page = int(item)
            if page < 1 or page > page_count:
                raise PdfToolkitError(f"Order page '{item}' is out of bounds.")
            result.append(page - 1)

        if not result:
            raise PdfToolkitError("No valid pages in page order.")
        return result
